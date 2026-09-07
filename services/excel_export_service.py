import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F2937"
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True
)

THIN_BORDER = Border(
    bottom=Side(
        style="thin",
        color="E5E7EB"
    )
)


def adjust_columns(worksheet):
    for column_cells in worksheet.columns:
        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            if cell.value is not None:
                length = len(str(cell.value))

                if length > max_length:
                    max_length = length

        worksheet.column_dimensions[
            column_letter
        ].width = min(max_length + 3, 40)


def style_worksheet(worksheet):
    worksheet.freeze_panes = "A2"

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    for row in worksheet.iter_rows(
        min_row=2
    ):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                vertical="top"
            )

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    adjust_columns(worksheet)


def create_excel_report(
    file_path,
    sheet_name,
    headers,
    rows
):
    os.makedirs(
        os.path.dirname(file_path),
        exist_ok=True
    )

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = sheet_name

    worksheet.append(headers)

    for row in rows:
        worksheet.append(
            list(row)
        )

    style_worksheet(
        worksheet
    )

    workbook.save(
        file_path
    )

    return file_path